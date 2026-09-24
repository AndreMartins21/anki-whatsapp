"""M14 (ADR-0017): dois alunos no privado, isolamento total. Cada chat é um espaço: palavras,
sessão, /list, /export e lembretes de um nunca aparecem no outro, e as travas e o limite de
"3 mensagens seguidas" são por espaço."""

from __future__ import annotations

import asyncio
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.domain.models import Estado, Profile
from app.services.fake_llm import FakeTutor
from app.services.lembretes import Agendador
from app.services.planilha import ExportadorExcel
from app.services.storage import ArmazenamentoEmMemoria
from tests.helpers import CHAT, T0, Montagem, expansoes, explicacao_stall, montar

CHAT_B = "5511988887777@c.us"


def _explicacao(palavra: str):  # type: ignore[no-untyped-def]
    return explicacao_stall().model_copy(update={"palavra": palavra})


async def _diz(m: Montagem, chat: str, texto: str) -> list[str]:
    antes = len(m.channel.textos_enviados)
    await m.router.processar(texto, chat)
    return [t for c, t in m.channel.textos_enviados[antes:] if c == chat]


def _dois_alunos(**kwargs: object) -> Montagem:
    tutor = FakeTutor(
        explicacoes=[_explicacao("stall"), _explicacao("hedge")],
        expansoes=[expansoes(), expansoes()],
    )
    return montar(tutor=tutor, **kwargs)  # type: ignore[arg-type]


async def test_a_sessao_de_um_aluno_nao_aparece_no_outro() -> None:
    m = _dois_alunos()

    await _diz(m, CHAT, "stall")

    assert m.banco.do_espaco(CHAT).obter_sessao().estado == Estado.AWAIT_ACTION
    assert m.banco.do_espaco(CHAT_B).obter_sessao().estado == Estado.IDLE


async def test_palavras_e_list_sao_de_cada_aluno() -> None:
    m = _dois_alunos()
    await _diz(m, CHAT, "stall")
    await _diz(m, CHAT, "3")  # salva
    await _diz(m, CHAT_B, "hedge")
    await _diz(m, CHAT_B, "3")

    lista_a = "\n".join(await _diz(m, CHAT, "/list"))
    lista_b = "\n".join(await _diz(m, CHAT_B, "/list"))

    assert "stall" in lista_a
    assert "hedge" not in lista_a
    assert "hedge" in lista_b
    assert "stall" not in lista_b
    assert [e.slug for e in m.banco.do_espaco(CHAT).listar_entradas()] == ["stall"]
    assert [e.slug for e in m.banco.do_espaco(CHAT_B).listar_entradas()] == ["hedge"]


async def test_aluno_sem_palavras_nao_ve_as_do_outro() -> None:
    m = _dois_alunos()
    await _diz(m, CHAT, "stall")
    await _diz(m, CHAT, "3")

    resposta = "\n".join(await _diz(m, CHAT_B, "/list"))

    assert "stall" not in resposta
    assert "don't have any saved words" in resposta


async def test_o_mesmo_slug_pode_existir_nos_dois_espacos() -> None:
    m = montar(
        tutor=FakeTutor(
            explicacoes=[_explicacao("stall"), _explicacao("stall")],
            expansoes=[expansoes(), expansoes()],
        )
    )
    await _diz(m, CHAT, "stall")
    await _diz(m, CHAT, "3")

    await _diz(m, CHAT_B, "stall")
    resposta = await _diz(m, CHAT_B, "3")

    assert resposta  # o B salvou sem esbarrar na palavra do A
    assert len(m.banco.do_espaco(CHAT_B).listar_entradas()) == 1


async def test_export_traz_so_as_palavras_de_quem_pediu() -> None:
    armazenamento = ArmazenamentoEmMemoria()
    m = _dois_alunos(exportador=ExportadorExcel(armazenamento, agora=lambda: T0))
    await _diz(m, CHAT, "stall")
    await _diz(m, CHAT, "3")
    await _diz(m, CHAT_B, "hedge")
    await _diz(m, CHAT_B, "3")

    await _diz(m, CHAT_B, "/export")

    ((chat, _nome, conteudo, _legenda),) = m.channel.arquivos_enviados
    assert chat == CHAT_B
    livro = load_workbook(BytesIO(conteudo))
    palavras = [r[1] for r in livro["Words"].iter_rows(min_row=2, values_only=True)]
    assert palavras == ["hedge"]


async def test_lembretes_sao_de_cada_espaco_e_so_o_dono_do_lembrete_recebe() -> None:
    m = montar()
    a, b = m.banco.do_espaco(CHAT), m.banco.do_espaco(CHAT_B)
    from tests.test_agendador import _entrada_vencida

    for espaco in (a, b):
        espaco.criar_entrada(_entrada_vencida())
    a.salvar_perfil(Profile(nivel="B1-B2", chat_id=CHAT, lembretes_por_dia=3, proximo_lembrete=T0))
    b.salvar_perfil(Profile(nivel="B1-B2", chat_id=CHAT_B, lembretes_por_dia=0))

    async def dormir(_: float) -> None:
        return None

    agendador = Agendador(
        router=m.router, banco=m.banco, agora=m.relogio.agora, fuso=ZoneInfo("UTC"), dormir=dormir
    )
    await agendador._tick()

    assert [c for c, _ in m.channel.textos_enviados] == [CHAT]
    perfil_b = b.obter_perfil()
    assert perfil_b is not None and perfil_b.lembrete_sem_resposta is False


async def test_lembrete_sem_resposta_de_um_espaco_nao_segura_o_outro() -> None:
    m = montar()
    from tests.test_agendador import _entrada_vencida

    for chat in (CHAT, CHAT_B):
        espaco = m.banco.do_espaco(chat)
        espaco.criar_entrada(_entrada_vencida())
        espaco.salvar_perfil(
            Profile(
                nivel="B1-B2",
                chat_id=chat,
                lembretes_por_dia=3,
                proximo_lembrete=T0,
                lembrete_sem_resposta=(chat == CHAT),  # o A não respondeu ao lembrete anterior
            )
        )

    async def dormir(_: float) -> None:
        return None

    agendador = Agendador(
        router=m.router, banco=m.banco, agora=m.relogio.agora, fuso=ZoneInfo("UTC"), dormir=dormir
    )
    await agendador._tick()

    assert [c for c, _ in m.channel.textos_enviados] == [CHAT_B]


async def test_falha_num_espaco_nao_impede_o_lembrete_do_outro() -> None:
    m = montar()
    from tests.test_agendador import _entrada_vencida

    m.banco.do_espaco(CHAT_B).criar_entrada(_entrada_vencida())
    m.banco.do_espaco(CHAT).salvar_perfil(
        Profile(nivel="B1-B2", chat_id=CHAT, lembretes_por_dia=3, proximo_lembrete=T0)
    )
    m.banco.do_espaco(CHAT_B).salvar_perfil(
        Profile(nivel="B1-B2", chat_id=CHAT_B, lembretes_por_dia=3, proximo_lembrete=T0)
    )

    class Quebrado:
        def __init__(self, banco: object) -> None:
            self._banco = banco

        def do_espaco(self, espaco_id: str) -> object:
            if espaco_id == CHAT:
                raise RuntimeError("Firestore fora do ar (só neste espaço)")
            return self._banco.do_espaco(espaco_id)  # type: ignore[attr-defined]

        def __getattr__(self, nome: str) -> object:
            return getattr(self._banco, nome)

    async def dormir(_: float) -> None:
        return None

    agendador = Agendador(
        router=m.router,
        banco=Quebrado(m.banco),  # type: ignore[arg-type]
        agora=m.relogio.agora,
        fuso=ZoneInfo("UTC"),
        dormir=dormir,
    )
    await agendador._tick()

    assert [c for c, _ in m.channel.textos_enviados] == [CHAT_B]


async def test_o_limite_de_tres_mensagens_seguidas_e_por_espaco() -> None:
    m = montar()
    for _ in range(3):
        await m.router.conversa_do(CHAT).enviar("oi")
    antes = len(m.channel.textos_enviados)

    await m.router.conversa_do(CHAT).enviar("quarta: descartada")
    await m.router.conversa_do(CHAT_B).enviar("primeira do outro aluno")

    novos = m.channel.textos_enviados[antes:]
    assert novos == [(CHAT_B, "primeira do outro aluno")]


async def test_chats_diferentes_andam_em_paralelo_e_o_mesmo_chat_um_por_vez() -> None:
    m = montar()
    liberar = asyncio.Event()
    entrou: list[str] = []

    async def presa(chat: str) -> None:
        async with m.router._trava(chat):
            entrou.append(chat)
            await liberar.wait()

    a1 = asyncio.create_task(presa(CHAT))
    b1 = asyncio.create_task(presa(CHAT_B))
    a2 = asyncio.create_task(presa(CHAT))
    await asyncio.sleep(0.01)

    # A e B entram juntos; a segunda tarefa do A espera a primeira terminar.
    assert sorted(entrou) == sorted([CHAT, CHAT_B])

    liberar.set()
    await asyncio.gather(a1, b1, a2)
    assert entrou.count(CHAT) == 2
