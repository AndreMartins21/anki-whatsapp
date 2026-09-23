"""Testes do export em Excel (seção 7.3, M12, ADR-0014): sem rede, com armazenamento em memória."""

from __future__ import annotations

from datetime import UTC, datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.workbook import Workbook

from app.domain.models import Entry, Sentence, SentidoSalvo, Synonym
from app.repo.memory import MemoryRepository
from app.services.planilha import TIPO_XLSX, ExportadorExcel, escolher_frase, gerar_planilha
from app.services.storage import ArmazenamentoEmMemoria

T0 = datetime(2026, 9, 20, 12, 0, tzinfo=UTC)


def _entrada(**campos: object) -> Entry:
    base: dict[str, object] = {
        "slug": "stall",
        "palavra": "stall",
        "classe": "verb",
        "cefr_estimado": "B2",
        "sentido": SentidoSalvo(traducao="travar", definicao="to stop making progress"),
        "outros_sentidos": [SentidoSalvo(traducao="enrolar", definicao="to delay")],
        "nota": "the car stalled",
        "tags": ["trabalho"],
        "origem_texto": "the talks stalled",
        "criado_em": T0,
        "atualizado_em": T0,
    }
    return Entry.model_validate(base | campos)


def _frase(texto: str, autor: str = "usuario", minutos: int = 0, **campos: object) -> Sentence:
    return Sentence.model_validate(
        {"texto": texto, "autor": autor, "criado_em": T0 + timedelta(minutes=minutos)} | campos
    )


def _abrir(conteudo: bytes) -> Workbook:
    return load_workbook(BytesIO(conteudo))


def test_tres_abas_com_cabecalho_e_uma_linha_por_palavra() -> None:
    entrada = _entrada(
        sinonimos=[
            Synonym(expressao="stumble", significado="to lose momentum", exemplo="It [[stumbled]].")
        ],
        repeticoes=2,
        lapsos=1,
        proxima_revisao=T0 + timedelta(days=3),
    )
    frases = [
        _frase("The talks [[stalled]].", autor="bot"),
        _frase(
            "the project stalled",
            minutos=5,
            veredito="correta",
            correcoes=["a → b", "c → d"],
            versao_natural="The project [[stalled]].",
            explicacao="Fine.",
        ),
    ]

    livro = _abrir(gerar_planilha([(entrada, frases)]))

    assert livro.sheetnames == ["Words", "Sentences", "Synonyms"]
    palavras = list(livro["Words"].iter_rows(values_only=True))
    assert palavras[0][:6] == ("#", "Word", "Class", "CEFR", "Translation", "Definition")
    linha = dict(zip(palavras[0], palavras[1], strict=True))
    assert linha["Word"] == "stall"
    assert linha["Other senses"] == "enrolar — to delay"
    assert linha["Tags"] == "trabalho"
    assert linha["Best sentence"] == "the project stalled"  # a do aluno, sem [[ ]]
    assert linha["Next review"] == datetime(2026, 9, 23, 12, 0)  # sem fuso
    assert (linha["Repetitions"], linha["Lapses"]) == (2, 1)

    sentencas = list(livro["Sentences"].iter_rows(min_row=2, values_only=True))
    assert [(s[2], s[3]) for s in sentencas] == [
        ("bot", "The talks stalled."),
        ("you", "the project stalled"),
    ]
    assert sentencas[1][4] == "The project stalled."
    assert sentencas[1][6] == "a → b\nc → d"

    sinonimos = list(livro["Synonyms"].iter_rows(min_row=2, values_only=True))
    assert sinonimos == [(1, "stall", "stumble", "to lose momentum", "It stumbled.")]


def test_cabecalho_em_negrito_e_congelado() -> None:
    livro = _abrir(gerar_planilha([(_entrada(), [])]))

    aba = livro["Words"]
    assert aba["A1"].font.bold is True
    assert aba.freeze_panes == "A2"


def test_escolher_frase_prefere_a_correta_do_aluno() -> None:
    frases = [
        _frase("The talks [[stalled]].", autor="bot"),
        _frase("I stalled the car", minutos=1, veredito="correta"),
    ]

    assert escolher_frase(_entrada(), frases) == "I [[stalled]] the car"


def test_escolher_frase_cai_na_versao_natural_e_depois_no_exemplo_do_bot() -> None:
    quase = _frase("x", minutos=1, veredito="quase", versao_natural="It [[stalled]] again.")
    exemplo = _frase("The talks [[stalled]].", autor="bot")

    assert escolher_frase(_entrada(), [exemplo, quase]) == "It [[stalled]] again."
    assert escolher_frase(_entrada(), [exemplo]) == "The talks [[stalled]]."
    assert escolher_frase(_entrada(), []) is None


def test_exportador_monta_a_planilha_sem_enviar_nada() -> None:
    repo = MemoryRepository()
    repo.criar_entrada(_entrada())
    repo.criar_entrada(_entrada(slug="hedge", palavra="hedge", criado_em=T0 + timedelta(minutes=1)))
    repo.adicionar_frase("stall", _frase("The talks [[stalled]].", autor="bot"))
    armazenamento = ArmazenamentoEmMemoria()

    resultado = ExportadorExcel(repo, armazenamento, agora=lambda: T0).exportar()

    assert resultado is not None
    assert resultado.quantidade == 2
    assert resultado.nome == "exports/vocabot_2026-09-20_1200.xlsx"
    assert armazenamento.arquivos == {}  # o bucket só entra no plano B
    livro = _abrir(resultado.conteudo)
    assert [r[1] for r in livro["Words"].iter_rows(min_row=2, values_only=True)] == [
        "stall",
        "hedge",
    ]


def test_plano_b_sobe_o_mesmo_arquivo_e_devolve_o_link() -> None:
    repo = MemoryRepository()
    repo.criar_entrada(_entrada())
    armazenamento = ArmazenamentoEmMemoria()
    resultado = ExportadorExcel(repo, armazenamento, agora=lambda: T0).exportar()
    assert resultado is not None

    link = resultado.gerar_link()

    assert link == "memoria://exports/vocabot_2026-09-20_1200.xlsx"
    assert armazenamento.arquivos["exports/vocabot_2026-09-20_1200.xlsx"] == resultado.conteudo


def test_exportador_sem_entradas_devolve_none() -> None:
    assert ExportadorExcel(MemoryRepository(), ArmazenamentoEmMemoria()).exportar() is None


def test_tipo_do_xlsx() -> None:
    assert TIPO_XLSX.endswith("spreadsheetml.sheet")
