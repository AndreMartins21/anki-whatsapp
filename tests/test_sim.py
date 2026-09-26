"""Testes do simulador de terminal (M6/M9): o ciclo completo, sem WhatsApp nem IA de verdade."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import load_workbook

from sim.__main__ import main


def _executar(falas: list[str], tmp_path: Path, *argv: str) -> tuple[int, str]:
    fila: Iterator[str] = iter(falas)
    saida: list[str] = []
    codigo = main(list(argv), entrada=lambda _: next(fila), saida=saida.append, exports=tmp_path)
    return codigo, "\n".join(saida)


def test_ciclo_completo_no_terminal(tmp_path: Path) -> None:
    codigo, tela = _executar(
        [
            "stall | the talks stalled",
            "2",  # see more examples
            "The negotiations stalled after the first meeting.",
            "4",  # just save
            "/list",
            "/info 1",
            "/export",
            "sair",
        ],
        tmp_path,
    )

    assert codigo == 0
    assert "*stall* (verb) — B2" in tela
    assert "📝 *Examples with stall*" in tela
    assert "✅ *Perfect!*" in tela
    assert "✅ Saved: *stall*." in tela
    assert "1. stall: travar, emperrar" in tela
    assert "✍️ *Your sentences*" in tela
    assert "1 word in the spreadsheet" in tela
    (arquivo,) = list(tmp_path.glob("vocabot_*.xlsx"))
    livro = load_workbook(arquivo)
    assert livro.sheetnames == ["Words", "Sentences", "Synonyms"]
    assert livro["Words"]["B2"].value == "stall"


def test_palavra_com_frase_de_contexto_mostra_o_card_direto(tmp_path: Path) -> None:
    _, tela = _executar(["stall | the talks stalled", "sair"], tmp_path)

    assert "*stall* (verb) — B2" in tela
    assert "Just save" in tela


def test_fim_da_entrada_encerra_sem_erro(tmp_path: Path) -> None:
    def acabou(_: str) -> str:
        raise EOFError

    codigo = main([], entrada=acabou, saida=lambda _: None, exports=tmp_path)

    assert codigo == 0


def test_real_llm_sem_configuracao_sai_com_erro(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    codigo = main(["--real-llm"], entrada=input, saida=lambda _: None, env={}, exports=tmp_path)

    assert codigo == 2
    assert "GCP_PROJECT_ID" in capsys.readouterr().err


def test_ciclo_de_musica_no_terminal(tmp_path: Path) -> None:
    """M13: as músicas do simulador são inventadas (sim/letras.py, ADR-0016)."""
    codigo, tela = _executar(
        [
            "/song paper plane",
            "1",  # The Inventors (há uma homônima)
            "ele deixou as chaves perto da porta",
            "no idea",  # o SimTutor aponta a palavra mais longa do verso
            "0",
            "all",
            "/song aviao de papel",
            "sair",
        ],
        tmp_path,
    )

    assert codigo == 0
    assert "2. *Paper Plane* — Quiet Harbor" in tela
    assert "Line 1/6" in tela
    assert "you went through 2 of 6 lines" in tela
    assert "Saved to your dictionary" in tela
    assert "aren't in English" in tela


# --- M16: `python -m sim --grupo` --------------------------------------------------------------


def test_ciclo_stall_no_grupo_com_dois_alunos(tmp_path: Path) -> None:
    codigo, tela = _executar(
        [
            "ana: !add stall | the talks stalled",
            "bia: !2",  # see more examples
            "ana: !The negotiations stalled after the first meeting.",
            "bia: !4",  # just save
            "ana: !list",
            "sair",
        ],
        tmp_path,
        "--grupo",
    )

    assert codigo == 0
    assert "*stall* (verb) — B2" in tela
    assert "!2 — See more examples" in tela
    assert "📝 *Examples with stall*" in tela
    assert "✅ *Perfect!*" in tela
    assert "✅ Saved: *stall*." in tela
    assert "The class's words" in tela and "1. stall: travar, emperrar" in tela


def test_no_grupo_linha_sem_prefixo_e_ignorada_e_nao_chega_ao_bot(tmp_path: Path) -> None:
    _, tela = _executar(
        ["ana: gente, alguém entendeu a aula?", "bia: !help", "sair"], tmp_path, "--grupo"
    )

    assert "(ignorado: sem prefixo, o bot não lê)" in tela
    assert tela.count("How I work in this group") == 1  # só o `!help` da Bia foi respondido


def test_no_grupo_linha_fora_do_formato_explica_o_uso(tmp_path: Path) -> None:
    _, tela = _executar(["!add stall", ": !help", "sair"], tmp_path, "--grupo")

    assert tela.count("(use `nome: mensagem`)") == 2


def test_no_grupo_o_dono_promove_um_professor_e_o_group_mostra_os_papeis(tmp_path: Path) -> None:
    _, tela = _executar(
        [
            "bia: !help",
            f"dono: !teacher {'5531990000001'}",
            "ana: !group",
            "sair",
        ],
        tmp_path,
        "--grupo",
        "--professor",
        "carla",
    )

    assert "• carla (teacher)" in tela
    assert "• bia (student)" in tela


def test_o_simulador_de_um_usuario_continua_igual(tmp_path: Path) -> None:
    _, tela = _executar(["stall", "sair"], tmp_path)

    assert "Simulador do vocabot" in tela and "grupo>" not in tela


# --- M17: revisão em grupo no simulador --------------------------------------------------------


def test_revisao_em_grupo_marca_o_aluno_pelo_nome_e_fecha_a_rodada(tmp_path: Path) -> None:
    _, tela = _executar(
        [
            "ana: !add stall | the talks stalled",
            "ana: !4",  # just save: a palavra vira um card da turma
            "carla: !help",  # a professora aparece, mas nunca é marcada
            "ana: !review",
            "ana: !it means to stop making progress",
            "sair",
        ],
        tmp_path,
        "--grupo",
        "--professor",
        "carla",
    )

    assert "@ana, your turn" in tela and "@carla" not in tela
    assert "Practice done" in tela


def test_timeout_no_simulador_repassa_o_card_e_depois_fecha(tmp_path: Path) -> None:
    _, tela = _executar(
        [
            "ana: !add stall",
            "ana: !4",
            "ana: !review",
            "~timeout",  # ninguém respondeu: repassa
            "~timeout",  # ninguém respondeu de novo: fecha
            "sair",
        ],
        tmp_path,
        "--grupo",
    )

    assert "No answer yet, so I'm passing this one on" in tela
    assert "Nobody answered in time" in tela
