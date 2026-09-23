"""Export em planilha Excel (seção 7.3 da spec, ADR-0014): um `.xlsx` com tudo o que o bot
coletou — palavras, frases (do aluno e do bot, com as correções) e sinônimos.

Substitui o arquivo de importação do Anki (M5 a M11). Três abas: `Words`, `Sentences`, `Synonyms`.
O arquivo vai direto pelo WhatsApp (ADR-0015); o bucket só entra como plano B.
"""

from __future__ import annotations

from collections.abc import Callable, Sequence
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.domain.choices import marcar_alvo
from app.domain.models import Entry, Sentence, agora_utc
from app.repo.base import Repository
from app.services.storage import Armazenamento

TIPO_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_LARGURA_MAXIMA = 60

_COLUNAS_PALAVRAS = (
    "#",
    "Word",
    "Class",
    "CEFR",
    "Translation",
    "Definition",
    "Other senses",
    "Note",
    "Tags",
    "Status",
    "Source text",
    "Best sentence",
    "Created",
    "Last review",
    "Next review",
    "Repetitions",
    "Lapses",
    "Ease",
)
_COLUNAS_FRASES = (
    "#",
    "Word",
    "Author",
    "Sentence",
    "Corrected version",
    "Verdict",
    "Corrections",
    "Explanation",
    "Date",
)
_COLUNAS_SINONIMOS = ("#", "Word", "Synonym", "Meaning", "Example")


@dataclass(frozen=True)
class ResultadoExportacao:
    """A planilha pronta para enviar. `gerar_link` é o plano B: sobe o mesmo arquivo no bucket e
    devolve a URL assinada — só é chamado se o envio direto pelo WhatsApp falhar."""

    nome: str
    conteudo: bytes
    quantidade: int
    gerar_link: Callable[[], str]


def _sem_marcas(texto: str | None) -> str:
    return (texto or "").replace("[[", "").replace("]]", "")


def _tem_marca(frase: str | None) -> bool:
    return frase is not None and "[[" in frase and "]]" in frase


def escolher_frase(entrada: Entry, frases: Sequence[Sentence]) -> str | None:
    """Melhor frase da palavra (seção 7.2), com o alvo entre [[ ]]: a melhor frase do usuário
    (`correta`, senão a `versao_natural` mais recente); sem frase do usuário, o primeiro exemplo
    do bot."""
    do_usuario = sorted((f for f in frases if f.autor == "usuario"), key=lambda f: f.criado_em)

    corretas = [f for f in do_usuario if f.veredito == "correta"]
    if corretas:
        melhor = corretas[-1]
        marcada = (
            melhor.texto if _tem_marca(melhor.texto) else marcar_alvo(melhor.texto, entrada.palavra)
        )
        if _tem_marca(marcada):
            return marcada
        if _tem_marca(melhor.versao_natural):
            return melhor.versao_natural

    for frase in reversed(do_usuario):
        if _tem_marca(frase.versao_natural):
            return frase.versao_natural

    for frase in sorted((f for f in frases if f.autor == "bot"), key=lambda f: f.criado_em):
        if _tem_marca(frase.texto):
            return frase.texto
    return None


def _data(valor: datetime | None) -> datetime | None:
    """O Excel não guarda fuso: converte para um horário "ingênuo" (UTC)."""
    return valor.replace(tzinfo=None) if valor else None


def _linha_da_palavra(n: int, e: Entry, frase: str | None) -> list[Any]:
    outros = "; ".join(f"{s.traducao} — {s.definicao}" for s in e.outros_sentidos)
    return [
        n,
        e.palavra,
        e.classe,
        e.cefr_estimado,
        e.sentido.traducao,
        e.sentido.definicao,
        outros,
        e.nota,
        ", ".join(e.tags),
        e.status,
        e.origem_texto or "",
        _sem_marcas(frase),
        _data(e.criado_em),
        _data(e.revisada_em),
        _data(e.proxima_revisao),
        e.repeticoes,
        e.lapsos,
        round(e.facilidade, 2),
    ]


def _linha_da_frase(n: int, e: Entry, f: Sentence) -> list[Any]:
    return [
        n,
        e.palavra,
        "you" if f.autor == "usuario" else "bot",
        _sem_marcas(f.texto),
        _sem_marcas(f.versao_natural),
        f.veredito or "",
        "\n".join(f.correcoes or []),
        f.explicacao or "",
        _data(f.criado_em),
    ]


def _preencher(aba: Worksheet, colunas: Sequence[str], linhas: Sequence[Sequence[Any]]) -> None:
    aba.append(list(colunas))
    for linha in linhas:
        aba.append(list(linha))
    for celula in aba[1]:
        celula.font = Font(bold=True)
    aba.freeze_panes = "A2"
    for indice, coluna in enumerate(aba.iter_cols(values_only=True), start=1):
        maior = max((len(str(v)) for v in coluna if v is not None), default=0)
        aba.column_dimensions[get_column_letter(indice)].width = min(maior + 2, _LARGURA_MAXIMA)


def gerar_planilha(entradas: Sequence[tuple[Entry, Sequence[Sentence]]]) -> bytes:
    """`.xlsx` com as três abas, na ordem em que as entradas foram passadas."""
    palavras: list[list[Any]] = []
    frases: list[list[Any]] = []
    sinonimos: list[list[Any]] = []
    for n, (entrada, frases_da_entrada) in enumerate(entradas, start=1):
        palavras.append(_linha_da_palavra(n, entrada, escolher_frase(entrada, frases_da_entrada)))
        for f in sorted(frases_da_entrada, key=lambda f: f.criado_em):
            frases.append(_linha_da_frase(len(frases) + 1, entrada, f))
        for s in entrada.sinonimos:
            sinonimos.append(
                [
                    len(sinonimos) + 1,
                    entrada.palavra,
                    s.expressao,
                    s.significado,
                    _sem_marcas(s.exemplo),
                ]
            )

    livro = Workbook()
    aba_palavras = livro.active
    assert aba_palavras is not None  # noqa: S101 — um Workbook novo sempre tem a aba ativa
    aba_palavras.title = "Words"
    _preencher(aba_palavras, _COLUNAS_PALAVRAS, palavras)
    _preencher(livro.create_sheet("Sentences"), _COLUNAS_FRASES, frases)
    _preencher(livro.create_sheet("Synonyms"), _COLUNAS_SINONIMOS, sinonimos)

    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


class ExportadorExcel:
    """Implementa o `Exportador` dos comandos: monta a planilha (não envia nada)."""

    def __init__(
        self,
        repo: Repository,
        armazenamento: Armazenamento,
        agora: Callable[[], datetime] = agora_utc,
    ) -> None:
        self._repo = repo
        self._armazenamento = armazenamento
        self._agora = agora

    def exportar(self) -> ResultadoExportacao | None:
        entradas = self._repo.listar_entradas()
        if not entradas:
            return None
        conteudo = gerar_planilha([(e, self._repo.listar_frases(e.slug)) for e in entradas])
        nome = f"exports/vocabot_{self._agora():%Y-%m-%d_%H%M}.xlsx"
        return ResultadoExportacao(
            nome=nome,
            conteudo=conteudo,
            quantidade=len(entradas),
            gerar_link=lambda: self._armazenamento.enviar(nome, conteudo, TIPO_XLSX),
        )
